"""Read source geometry and render with native PowerPoint, without saving sources."""
import argparse
import hashlib
import json
import ctypes
import time
from pathlib import Path

import win32com.client
from pptx import Presentation


def inspect(source, output):
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation(source)
    report = {"source": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
              "width_emu": prs.slide_width, "height_emu": prs.slide_height, "slides": []}
    scale = 1536 / prs.slide_width
    for index, slide in enumerate(prs.slides, 1):
        shapes = []
        for shape in slide.shapes:
            info = {"id": shape.shape_id, "name": shape.name,
                    "box": [round(v * scale, 3) for v in (shape.left, shape.top, shape.width, shape.height)]}
            if shape.has_text_frame:
                info["text"] = shape.text
                info["runs"] = []
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        font = run.font
                        color = str(font.color.rgb) if font.color.type and font.color.type == 1 else None
                        info["runs"].append({"text": run.text, "font": font.name,
                                             "pt": font.size.pt if font.size else None,
                                             "bold": font.bold, "color": color})
            shapes.append(info)
        report["slides"].append({"number": index, "shapes": shapes})
    (output / "geometry.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    # Register only for this Windows session; no registry or font-folder changes.
    fonts = list((Path(__file__).resolve().parents[1] / "assets" / "fonts").glob("*.ttf"))
    for font in fonts:
        if not ctypes.windll.gdi32.AddFontResourceExW(str(font), 0, None):
            raise RuntimeError(f"Could not register {font.name}")
    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.AutomationSecurity = 3
    presentation = None
    try:
        presentation = app.Presentations.Open(str(source.resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
        width = 3072
        height = round(width * prs.slide_height / prs.slide_width)
        for number in range(1, presentation.Slides.Count + 1):
            for attempt in range(3):
                try:
                    presentation.Slides(number).Export(str((output / f"slide-{number}.png").resolve()), "PNG", width, height)
                    break
                except Exception:
                    if attempt == 2:
                        raise
                    time.sleep(1)
        presentation.SaveAs(str((output / "reference.pdf").resolve()), 32)
        print(json.dumps({"source": source.name, "pages": presentation.Slides.Count,
                          "render_size": [width, height], "output": str(output)}, ensure_ascii=True))
    finally:
        if presentation:
            presentation.Close()
        app.Quit()
        for font in fonts:
            ctypes.windll.gdi32.RemoveFontResourceExW(str(font), 0, None)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    inspect(args.source, args.output)
